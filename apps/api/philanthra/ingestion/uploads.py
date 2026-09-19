"""Aggregate-only tabular onboarding and quarantined report extraction."""

import csv
import io
import multiprocessing
import re
import time
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from .common import (
    MAX_CELL,
    MAX_COLUMNS,
    MAX_PAGES,
    MAX_ROWS,
    MAX_SECONDS,
    ImportProblem,
    archive_entries,
    deadline_check,
    decimal_string,
    envelope,
    finish,
    issue,
    size_check,
    xml_root,
)

COMMON = {
    "record_type",
    "program_name",
    "cause",
    "intervention",
    "population",
    "geography",
    "period_start",
    "period_end",
}
OUTCOME = {
    "outcome_code",
    "outcome_definition",
    "unit",
    "direction",
    "numerator",
    "denominator",
    "cohort_id",
    "study_id",
    "independent",
    "followup_months",
    "comparator",
    "study_design",
    "uncertainty",
    "missingness",
    "stratum",
}
COST = {"amount", "currency", "unit", "denominator"}
FIELDS = COMMON | OUTCOME | COST
SENSITIVE = re.compile(
    r"(?:beneficiar|patient|person|participant|household)[ _-]?(?:id|name|address|email|phone)|"
    r"(?:^|[ _-])(?:ssn|email|phone|dob|birth|address|passport|diagnosis|ethnicity|religion|"
    r"race|gender|hiv|disability|first_name|last_name)(?:$|[ _-])",
    re.I,
)
FORMULA = re.compile(r"^[\s\x00-\x1f]*[=+@]|^[\s\x00-\x1f]*-(?!\d+(?:\.\d+)?\s*$)")
SUSPICIOUS = re.compile(
    r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b|\b\d{3}-\d{2}-\d{4}\b|"
    r"ignore (?:all |previous |prior )*instructions|<script\b",
    re.I,
)


def safe_csv(records, columns=None):
    """Neutralize spreadsheet formula cells, including leading whitespace/control bytes."""
    rows = list(records)
    columns = columns or list(dict.fromkeys(k for row in rows for k in row))
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)

    def cell(value):
        text = "" if value is None else str(value)
        if text.lstrip(" \t\r\n\x00").startswith(("=", "+", "-", "@")):
            text = "'" + text
        return text

    writer.writerow([cell(c) for c in columns])
    for row in rows:
        writer.writerow([cell(row.get(c)) for c in columns])
    return buffer.getvalue()


def rejected_csv(result):
    """Rejected export contains locations and safe issues, never raw sensitive rows."""
    return safe_csv(
        [d for d in result["diagnostics"] if d["severity"] == "error"],
        ["row", "field", "code", "message"],
    )


def _tabular(payload, suffix):
    if suffix == ".csv":
        try:
            text = payload.decode("utf-8-sig")
            if "\x00" in text:
                raise ImportProblem("invalid_text", "NUL bytes are not supported.")
            csv.field_size_limit(MAX_CELL)
            rows = []
            for row in csv.reader(io.StringIO(text), strict=True):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ImportProblem("row_limit", "The file exceeds 5,000 data rows.")
            return rows
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ImportProblem(
                "invalid_csv", "CSV must be well-formed UTF-8 with bounded fields."
            ) from exc
    entries = archive_entries(payload)
    if "[Content_Types].xml" not in entries or "xl/workbook.xml" not in entries:
        raise ImportProblem("invalid_xlsx", "The file is not an XLSX workbook.")
    if any("vba" in name.lower() or "externallinks" in name.lower() for name in entries):
        raise ImportProblem("unsafe_workbook", "Macros and external links are not supported.")
    for name, content in entries.items():
        if name.endswith(".rels"):
            relationships = xml_root(content)
            if any(node.get("TargetMode") == "External" for node in relationships):
                raise ImportProblem(
                    "unsafe_workbook", "External workbook relationships are prohibited."
                )
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
        if len(workbook.worksheets) != 1:
            raise ImportProblem("worksheet_count", "Provide exactly one worksheet.")
        sheet = workbook.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ImportProblem("row_limit", "The worksheet exceeds 5,000 data rows.")
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ImportProblem("column_limit", "The worksheet exceeds 64 columns.")
        rows = []
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                if cell.data_type == "f" or getattr(cell, "hyperlink", None):
                    raise ImportProblem(
                        "spreadsheet_formula", "Formulas and active hyperlinks are prohibited."
                    )
                value = cell.value
                if isinstance(value, datetime):
                    value = value.date().isoformat()
                elif isinstance(value, date):
                    value = value.isoformat()
                values.append("" if value is None else str(value))
            rows.append(values)
            if len(rows) > MAX_ROWS + 1:
                raise ImportProblem("row_limit", "The worksheet exceeds 5,000 data rows.")
        workbook.close()
        return rows
    except ImportProblem:
        raise
    except Exception as exc:
        raise ImportProblem("invalid_xlsx", "The workbook could not be parsed safely.") from exc


def _normalize(row, row_number, result):
    kind = row.get("record_type")
    if kind not in {"outcome", "cost"}:
        raise ImportProblem("record_type", "record_type must be outcome or cost.")
    required = COMMON | (
        {
            "outcome_code",
            "outcome_definition",
            "unit",
            "direction",
            "cohort_id",
            "study_id",
            "study_design",
            "followup_months",
            "independent",
        }
        if kind == "outcome"
        else {"amount", "currency", "unit"}
    )
    if any(not row.get(field) for field in required):
        missing = ", ".join(sorted(field for field in required if not row.get(field)))
        raise ImportProblem("required_field", "Required normalized fields missing: " + missing)
    for field in ("period_start", "period_end"):
        try:
            row[field] = date.fromisoformat(row[field]).isoformat()
        except (TypeError, ValueError) as exc:
            raise ImportProblem("invalid_date", "Dates must use YYYY-MM-DD.") from exc
    if row["period_start"] > row["period_end"]:
        raise ImportProblem("invalid_period", "The reporting end date precedes its start date.")
    for field in ("numerator", "denominator", "followup_months"):
        if field not in row:
            continue
        value = row[field]
        if value in (None, ""):
            row[field] = None
        elif not re.fullmatch(r"\d{1,9}", value):
            raise ImportProblem(
                "invalid_count", "Counts and follow-up months must be nonnegative integers."
            )
        else:
            row[field] = int(value)
    if kind == "outcome":
        if row["unit"] != "binary" or row["direction"] not in {"higher_better", "lower_better"}:
            raise ImportProblem(
                "unsupported_outcome", "Use a binary outcome with an explicit direction."
            )
        if row["independent"].lower() not in {"true", "false"}:
            raise ImportProblem(
                "independence", "Cohort independence must be explicitly true or false."
            )
        row["independent"] = row["independent"].lower() == "true"
        row.setdefault("numerator", None)
        row.setdefault("denominator", None)
        if row["denominator"] == 0:
            raise ImportProblem("invalid_denominator", "A reported denominator must be positive.")
        if row["numerator"] is not None and row["denominator"] is not None:
            if row["numerator"] > row["denominator"]:
                raise ImportProblem("invalid_numerator", "The numerator exceeds the denominator.")
        else:
            issue(
                result,
                "missing_denominator_or_numerator",
                "Incomplete outcome cannot be pooled.",
                row=row_number,
                severity="warning",
            )
        if row["study_design"] not in {
            "observational",
            "correlational",
            "quasi_experimental",
            "randomized",
        }:
            raise ImportProblem("study_design", "Unsupported study design.")
        if row["study_design"] in {"quasi_experimental", "randomized"} and not row.get(
            "comparator"
        ):
            raise ImportProblem("missing_comparator", "This study design requires a comparator.")
    else:
        row["amount"] = decimal_string(row["amount"], nonnegative=True)
        exponent = Decimal(row["amount"]).as_tuple().exponent
        if not isinstance(exponent, int) or exponent < -2:
            raise ImportProblem("money_precision", "Costs support at most two fractional digits.")
        if row["currency"] not in {
            "USD",
            "EUR",
            "GBP",
            "MXN",
            "ETB",
            "GTQ",
            "CAD",
            "AUD",
            "JPY",
            "CHF",
        }:
            raise ImportProblem(
                "currency", "Currency is not in the pilot's supported ISO currency list."
            )
        if row.get("denominator") == 0:
            raise ImportProblem("invalid_denominator", "A reported denominator must be positive.")
    allowed = COMMON | (OUTCOME if kind == "outcome" else COST)
    if any(row.get(field) not in (None, "") for field in set(row) - allowed):
        raise ImportProblem("incompatible_field", "A value belongs to another record type.")
    row = {key: value for key, value in row.items() if key in allowed}
    row["locator"] = f"row:{row_number}"
    return row


def _document_child(payload, suffix, pipe):
    try:
        import resource
        import sys

        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        if sys.platform == "linux":
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        if suffix == ".txt":
            pages = [payload.decode("utf-8-sig")]
        else:
            from pypdf import PdfReader, overwrite_configuration

            overwrite_configuration(
                maximum_declared_stream_length=2 * 1024 * 1024,
                array_based_stream_maximum_output_length=2 * 1024 * 1024,
                zlib_maximum_output_length=2 * 1024 * 1024,
                lzw_maximum_output_length=2 * 1024 * 1024,
                run_length_maximum_output_length=2 * 1024 * 1024,
                image_maximum_buffer_size=2 * 1024 * 1024,
            )
            reader = PdfReader(io.BytesIO(payload), strict=True)
            if reader.is_encrypted or len(reader.pages) > MAX_PAGES:
                raise ImportProblem(
                    "pdf_limit", "Encrypted PDFs or more than 40 pages are unsupported."
                )
            pages = []
            for page in reader.pages:
                content = page.get_contents()
                if content and len(content.get_data()) > 2 * 1024 * 1024:
                    raise ImportProblem("pdf_stream_limit", "PDF content stream exceeds the limit.")
                pages.append(page.extract_text() or "")
        if sum(len(p) for p in pages) > 200000:
            raise ImportProblem("text_limit", "Extracted report text exceeds 200,000 characters.")
        pipe.send({"pages": pages})
    except ImportProblem as exc:
        pipe.send({"code": exc.code, "message": str(exc)})
    except Exception:
        pipe.send(
            {
                "code": "text_unavailable",
                "message": "Text extraction failed; provide text or aggregate CSV.",
            }
        )
    finally:
        pipe.close()


def _document(payload, suffix, result):
    if suffix == ".pdf" and not payload.startswith(b"%PDF-"):
        raise ImportProblem("invalid_pdf", "The file does not have a PDF signature.")
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_document_child, args=(payload, suffix, send), daemon=True)
    process.start()
    send.close()
    try:
        if not receive.poll(MAX_SECONDS):
            raise ImportProblem(
                "time_limit", "Document extraction exceeded its isolated processing limit."
            )
        try:
            extracted = receive.recv()
        except EOFError as exc:
            raise ImportProblem(
                "text_unavailable", "Document extraction stopped at its resource limit."
            ) from exc
    finally:
        if process.is_alive():
            process.terminate()
        process.join(timeout=2)
        receive.close()
    if "code" in extracted:
        raise ImportProblem(extracted["code"], extracted["message"])
    pages = extracted["pages"]
    result["status"] = "quarantined"
    if not any(page.strip() for page in pages):
        issue(
            result,
            "text_unavailable",
            "No text extracted; provide text or aggregate CSV.",
            severity="warning",
        )
    elif any(SUSPICIOUS.search(page) for page in pages):
        issue(
            result,
            "sensitive_or_instruction_content",
            "Report requires restricted manual review; extracted content is withheld.",
            severity="warning",
        )
    else:
        result["preview"] = [
            {"page": i, "locator": f"page:{i}", "text": page} for i, page in enumerate(pages, 1)
        ]
    issue(
        result,
        "manual_review",
        "Report text is quarantined and is not structured evidence or analytics input.",
        severity="warning",
    )
    return result


def parse_upload(payload: bytes, filename: str, mapping: dict | None = None) -> dict:
    suffix = Path(filename).suffix.lower()
    result = envelope(payload, "document" if suffix in {".txt", ".pdf"} else "aggregate")
    start = time.monotonic()
    try:
        size_check(payload)
        if suffix in {".pdf", ".txt"}:
            return _document(payload, suffix, result)
        if suffix not in {".csv", ".xlsx"}:
            raise ImportProblem(
                "unsupported_type", "Supported uploads are CSV, XLSX, text PDF and UTF-8 text."
            )
        rows = _tabular(payload, suffix)
        if not rows or not rows[0] or len(rows) < 2:
            raise ImportProblem("empty_table", "Provide a header and at least one data row.")
        columns = [str(c).strip() for c in rows[0]]
        if len(columns) > MAX_COLUMNS or any(not c or len(c) > 100 for c in columns):
            raise ImportProblem(
                "invalid_columns",
                "Headers must be nonempty, at most 100 characters and at most 64 columns.",
            )
        if len(set(columns)) != len(columns):
            raise ImportProblem("duplicate_column", "Column names must be unique.")
        if any(SENSITIVE.search(c) or FORMULA.search(c) or SUSPICIOUS.search(c) for c in columns):
            raise ImportProblem(
                "sensitive_column",
                "Beneficiary identifiers and sensitive or active column names are prohibited.",
            )
        result["columns"] = columns
        if mapping is None:
            result["status"] = "needs_mapping"
            return result
        if not isinstance(mapping, dict) or set(mapping) != set(columns):
            raise ImportProblem(
                "unmapped_columns",
                "Explicitly map every column; remove unsupported columns before uploading.",
            )
        if any(not isinstance(v, str) or v not in FIELDS for v in mapping.values()):
            raise ImportProblem(
                "unsupported_mapping", "Every column must map to a supported aggregate field."
            )
        if len(set(mapping.values())) != len(mapping):
            raise ImportProblem(
                "duplicate_mapping", "Each normalized field may be mapped only once."
            )
        seen = set()
        for number, values in enumerate(rows[1:], 2):
            deadline_check(start)
            if not any(values):
                continue
            try:
                if len(values) != len(columns):
                    raise ImportProblem("column_count", "The row width does not match its header.")
                if any(len(value) > MAX_CELL for value in values):
                    raise ImportProblem("cell_limit", "A cell exceeds the text limit.")
                if any(FORMULA.search(value) for value in values):
                    raise ImportProblem(
                        "spreadsheet_formula", "Formula-like cell content is prohibited."
                    )
                if any(SUSPICIOUS.search(value) for value in values):
                    raise ImportProblem(
                        "sensitive_content",
                        "Potential identifiers or instructions require restricted manual review.",
                    )
                row = _normalize(
                    {mapping[c]: value.strip() for c, value in zip(columns, values, strict=True)},
                    number,
                    result,
                )
                identity = tuple(
                    (key, str(value)) for key, value in sorted(row.items()) if key != "locator"
                )
                if identity in seen:
                    raise ImportProblem(
                        "duplicate_row", "This row duplicates an earlier normalized record."
                    )
                seen.add(identity)
                result["records"].append(row)
            except ImportProblem as exc:
                issue(result, exc.code, str(exc), row=number)
        result["preview"] = result["records"][:10]
    except ImportProblem as exc:
        issue(result, exc.code, str(exc))
    return finish(result)
